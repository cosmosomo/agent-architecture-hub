# -*- coding: utf-8 -*-
"""把 notes.json 写进 pptx 的讲者备注，并为第 4 页的原始全尺寸截图设置 PowerPoint 原生无损裁剪（crop_left/top/right/bottom）。
用户在 PowerPoint 中右键任意截图点击「裁剪」，即可直接拖拽调整完整原图的视野范围！
"""
import json
import sys

from lxml import etree
from pptx import Presentation

BODY = ('<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        '<p:nvSpPr><p:cNvPr id="%d" name="Notes Placeholder"/><p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr>'
        '<p:nvPr><p:ph type="body" idx="1"/></p:nvPr></p:nvSpPr>'
        '<p:spPr><a:xfrm><a:off x="685800" y="6096000"/><a:ext cx="5486400" cy="5486400"/></a:xfrm></p:spPr>'
        '<p:txBody><a:bodyPr/><a:lstStyle/><a:p/></p:txBody></p:sp>')

deck, notes_path = sys.argv[1], sys.argv[2]
notes = {k: v for k, v in json.load(open(notes_path, encoding='utf-8')).items() if not k.startswith('_')}
prs = Presentation(deck)

# 1. 写入讲者备注
for k, text in sorted(notes.items(), key=lambda kv: int(kv[0])):
    ns = prs.slides[int(k) - 1].notes_slide
    if ns.notes_text_frame is None:
        tree = ns.shapes._spTree
        ids = [int(i) for i in tree.xpath('.//p:cNvPr/@id')] or [1]
        tree.append(etree.fromstring(BODY % (max(ids) + 1)))
    ns.notes_text_frame.text = text

# 2. 为第 4 页的独立全尺寸原图应用 PowerPoint 原生无损局部裁剪 (按画布从左到右精确匹配)
if len(prs.slides) >= 4:
    s4 = prs.slides[3]
    # 仅筛选内容区图片 (100pt < top < 300pt)，按 (left, top) 从左到右排序，绝不触碰顶部校徽与底图
    content_pics = sorted(
        [sh for sh in s4.shapes if sh.shape_type == 13 and 100 < sh.top.pt < 300],
        key=lambda s: (round(s.left.pt, 1), round(s.top.pt, 1)),
    )
    # 排序后严格对应:
    # 0 (left=47.2):  orig_blackboard_light.png (1920x1080) -> 聚焦中央黑板引力场与角色四分外环
    # 1 (left=51.8):  orig_ring_badge.png (322x150)         -> 左下独立悬浮「送达四分外环」卡片（可独立拖动）
    # 2 (left=342.8): orig_console_light.png (1280x800)     -> 聚焦「② 个体名册与能力渠道负载矩阵」
    # 3 (left=491.2): orig_reverse_diff.png (1920x1080)     -> 聚焦右侧「⚡ 反向调控闭环 + inbox.md Diff」
    # 4 (left=636.8): orig_m1_feed.png (375x812)            -> 聚焦手机端「待处理流 (22 待办)」上半屏
    # 5 (left=727.5): orig_m2_steer.png (375x812)           -> 聚焦手机端「实时会话 + 底部 steer 插话发送框」下半屏
    # 6 (left=818.2): orig_m3_projects.png (375x812)        -> 聚焦手机端「跨家活跃项目池」上半屏
    crops = [
        (535 / 1920, 86 / 1080, (1920 - 1415) / 1920, (1080 - 728) / 1080),
        (0.0, 0.0, 0.0, 0.0),
        (748 / 1280, 188 / 800, (1280 - 1226) / 1280, (800 - 584) / 800),
        (1546 / 1920, 360 / 1080, (1920 - 1914) / 1920, (1080 - 836) / 1080),
        (0.0, 0.0, 0.0, (812 - 560) / 812),
        (0.0, (812 - 560) / 812, 0.0, 0.0),
        (0.0, 0.0, 0.0, (812 - 560) / 812),
    ]
    for sh, (cl, ct, cr, cb) in zip(content_pics, crops):
        sh.crop_left = round(cl, 5)
        sh.crop_top = round(ct, 5)
        sh.crop_right = round(cr, 5)
        sh.crop_bottom = round(cb, 5)
    print(f'applied native PowerPoint non-destructive crops to {len(content_pics)} content pictures on Slide 4')

prs.save(deck)
print('notes ->', ', '.join(sorted(notes, key=int)))
